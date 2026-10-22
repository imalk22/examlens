"""Excel and CSV export helpers."""

from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="2563EB")
LEVEL_FILLS = {
    "High": PatternFill("solid", fgColor="FECACA"),
    "Medium": PatternFill("solid", fgColor="FDE68A"),
    "Low": PatternFill("solid", fgColor="BBF7D0"),
}
RESULT_FILLS = {
    "Pass": PatternFill("solid", fgColor="BBF7D0"),
    "Fail": PatternFill("solid", fgColor="FECACA"),
    "Absent": PatternFill("solid", fgColor="E5E7EB"),
}


def to_csv_bytes(df: pd.DataFrame) -> bytes:
    """UTF-8 CSV (with BOM so Excel opens it correctly)."""
    return df.to_csv(index=False).encode("utf-8-sig")


def _prepare(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for col in out.columns:
        if pd.api.types.is_datetime64_any_dtype(out[col]):
            out[col] = out[col].dt.date
    return out


def _format_sheet(ws, df: pd.DataFrame, color_col: str | None, fills: dict) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for idx, col in enumerate(df.columns, start=1):
        longest = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 10), 70)
    ws.freeze_panes = "A2"
    if len(df):
        ws.auto_filter.ref = ws.dimensions
    if color_col and color_col in df.columns:
        col_idx = list(df.columns).index(color_col) + 1
        for row in range(2, len(df) + 2):
            cell = ws.cell(row=row, column=col_idx)
            if cell.value in fills:
                cell.fill = fills[cell.value]


def to_excel_bytes(
    sheets: dict[str, pd.DataFrame],
    color_col: str | None = None,
    fills: dict | None = None,
) -> bytes:
    """Write each DataFrame to its own formatted sheet and return the .xlsx bytes."""
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, df in sheets.items():
            prepared = _prepare(df)
            prepared.to_excel(writer, sheet_name=name[:31], index=False)
            _format_sheet(writer.sheets[name[:31]], prepared, color_col, fills or {})
    return buffer.getvalue()


def at_risk_workbook(risk: pd.DataFrame) -> bytes:
    """Coordinator-ready Excel of the at-risk list, coloured by risk level."""
    cols = {
        "student_id": "Student ID",
        "student_name": "Name",
        "programme": "Programme",
        "intake": "Intake",
        "risk_level": "Risk level",
        "risk_score": "Risk score",
        "outstanding_subjects": "Subjects outstanding",
        "latest_avg": "Latest session average",
        "reasons": "Reasons",
    }
    table = risk[list(cols)].rename(columns=cols).round({"Latest session average": 1})
    return to_excel_bytes({"At-risk students": table}, "Risk level", LEVEL_FILLS)


def transcript_workbook(
    info: dict[str, object], history: pd.DataFrame, remaining: pd.DataFrame
) -> bytes:
    """Transcript-style workbook: summary, full exam history and outstanding subjects."""
    summary = pd.DataFrame({"Field": list(info), "Value": [str(v) for v in info.values()]})
    cols = [
        "exam_session",
        "subject_code",
        "subject_name",
        "attempt_no",
        "marks",
        "grade",
        "result",
    ]
    hist = history.sort_values(["session_date", "subject_code"])[cols]
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame, color in [
            ("Summary", summary, None),
            ("Exam history", hist, "result"),
            ("Outstanding", remaining, None),
        ]:
            prepared = _prepare(frame)
            prepared.to_excel(writer, sheet_name=name, index=False)
            _format_sheet(writer.sheets[name], prepared, color, RESULT_FILLS)
    return buffer.getvalue()
